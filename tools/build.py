# -*- coding: utf-8 -*-
"""Bước 5 — đóng gói giáo trình Word."""
import sys, os, re, glob
HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.environ.get('HSBG_WORK', os.path.join(HERE, '..', 'build'))
REPO = os.path.abspath(os.path.join(HERE, '..'))
sys.path.insert(0, HERE)
import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from style import *
from style import _el
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from render import Renderer, emit_runs

DRAFTS = os.path.join(REPO, 'drafts')
OUT = os.path.join(REPO, 'Ho_so_bai_giang_lai_xe_2026_v1.0.docx')
VERSION = 'Bản 1.0'
UPDATED = 'Tháng 9/2026'

doc = Document()

# ═══════════════ THIẾT LẬP TRANG & STYLE ═══════════════
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.0)

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(10.5)
st.font.color.rgb = RGBColor.from_string(INK)
st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.18

for nm, sz, col in (('Heading 1', 19, NAVY), ('Heading 2', 13.5, BLUE),
                    ('Heading 3', 11.5, BLUE_MID)):
    s = doc.styles[nm]
    s.font.name = FONT; s.font.size = Pt(sz); s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(col)
    s._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    s.paragraph_format.keep_with_next = True
for nm in ('List Bullet', 'List Number'):
    s = doc.styles[nm]
    s.font.name = FONT; s.font.size = Pt(10.5)
    s.font.color.rgb = RGBColor.from_string(INK)

# ═══════════════ TRANG BÌA ═══════════════
def cover():
    # dải màu trên
    t = doc.add_table(rows=1, cols=1); no_borders(t); cell_margins(t, 0, 0, 0, 0)
    c = t.cell(0, 0); shade(c, NAVY); c.width = Cm(16.5)
    p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(' '); r.font.size = Pt(5)

    for _ in range(2):
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
        sp.add_run().font.size = Pt(11)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('CƠ SỞ ĐÀO TẠO LÁI XE')
    r.font.size = Pt(11); r.bold = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(GREY_TXT)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run('…………………………………………')
    r.font.size = Pt(11); r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(GREY_LINE)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('HỒ SƠ BÀI GIẢNG')
    r.font.size = Pt(34); r.bold = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run('THỰC HÀNH LÁI XE Ô TÔ')
    r.font.size = Pt(21); r.bold = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(BLUE)

    # thanh hạng đào tạo
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t); cell_margins(t, 130, 200, 130, 200)
    c = t.cell(0, 0); shade(c, BLUE_PALE); c.width = Cm(11)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run('HẠNG B  ·  HẠNG C1')
    r.font.size = Pt(16); r.bold = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(3); p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run('Xe số tự động và xe số cơ khí (số sàn)')
    r.font.size = Pt(10); r.italic = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(BLUE)

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(18)

    # ảnh bìa: xe tập lái
    import os
    from render import MEDIA, EXTENTS
    fn = 'image3.jpeg'
    path = os.path.join(MEDIA, fn)
    if os.path.exists(path):
        cx, cy = EXTENTS.get(fn, (Inches(4), Inches(2.5)))
        target = Inches(3.9)
        cy = int(cy * target / cx); cx = int(target)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        add_picture_any(doc, p, path, cx, cy)

    # khối thông tin
    t = doc.add_table(rows=4, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t); cell_margins(t, 50, 60, 50, 140)
    info = [('Môn học', 'Thực hành lái xe ô tô'),
            ('Đối tượng', 'Học viên học lái xe hạng B và hạng C1'),
            ('Biên soạn', '……………………………………………'),
            ('Cập nhật', '%s — %s' % (UPDATED, VERSION))]
    for i, (k, v) in enumerate(info):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        c0.width, c1.width = Cm(3.2), Cm(9.0)
        p = c0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(k); r.bold = True; r.font.size = Pt(10); r.font.name = FONT
        r.font.color.rgb = RGBColor.from_string(GREY_TXT)
        p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
        r = p.add_run(v); r.font.size = Pt(10); r.font.name = FONT
        r.font.color.rgb = RGBColor.from_string(INK)

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('Biên soạn theo Luật 36/2024/QH15 · Thông tư 17/2026/TT-BXD · Thông tư 108/2026/TT-BCA')
    r.font.size = Pt(8.5); r.italic = True; r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(GREY_TXT)

cover()

# ═══════════════ TRANG HƯỚNG DẪN SỬ DỤNG ═══════════════
p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
rnd = Renderer(doc)

GUIDE = """
## HƯỚNG DẪN SỬ DỤNG TÀI LIỆU

Tài liệu này là **cẩm nang học lái xe** dùng cho cả học viên và giáo viên. Mỗi bài học được trình bày theo **một cấu trúc thống nhất gồm 8 phần**, giúp bạn tìm nhanh nội dung cần thiết ở bất kỳ thời điểm nào của khoá học.

| Phần | Trả lời câu hỏi |
|:---:|---|
| 01. Mục tiêu bài học | Học xong bài này tôi làm được gì? |
| 02. Chuẩn bị | Cần xe gì, sân gì, điều kiện an toàn nào? |
| 03. Nội dung bài học | Thông số xe, kích thước hình tập, điểm chuẩn ở đâu? |
| 04. Quy trình thực hiện | Tôi phải làm những bước nào, theo thứ tự ra sao? |
| 05. Điểm cần lưu ý | Tôi cần chú ý điều gì để không mất an toàn và không mất điểm? |
| 06. Lỗi thường gặp | Tôi thường sai ở đâu và sửa thế nào? |
| 07. Tiêu chí hoàn thành | Thế nào là đạt? |
| 08. Phiếu tự kiểm tra | Tôi đã đạt yêu cầu chưa? |

## HỆ THỐNG HỘP THÔNG TIN

Trong tài liệu có 6 loại hộp thông tin, mỗi loại một màu và một biểu tượng riêng:

| Hộp | Nội dung bên trong |
|---|---|
| 📖 **GHI NHỚ** | Thông tin học viên cần thuộc lòng |
| ⚠️ **LƯU Ý** | Điều dễ nhầm lẫn, cần đọc kỹ |
| ⚠️ **LƯU Ý AN TOÀN** | Nội dung liên quan đến an toàn người và phương tiện |
| 💡 **MẸO QUAN SÁT** | Cách quan sát, căn điểm chuẩn, định hướng |
| 📋 **KIỂM TRA** | Đối chiếu với tiêu chí chấm điểm ở kỳ sát hạch |
| ⚠️ **LƯU Ý CHUYỂN TIẾP** | Thay đổi pháp lý theo mốc thời gian — **rất quan trọng** |

## HỆ THỐNG NHÃN TRUY VẾT NGUỒN

Tài liệu này được cập nhật từ hồ sơ bài giảng cũ. Để bạn phân biệt rõ đâu là nội dung gốc, đâu là nội dung cập nhật, mỗi thay đổi đều được gắn nhãn:

| Nhãn | Ý nghĩa |
|---|---|
| *(không nhãn)* | Nội dung từ hồ sơ bài giảng gốc, giữ nguyên |
| `[CẬP NHẬT 2026]` | Nội dung được thay theo văn bản pháp luật mới, có ghi căn cứ |
| `[BIÊN SOẠN MỚI]` | Nội dung soạn thêm, không có trong hồ sơ gốc |
| `[CẦN RÀ SOÁT]` | Nội dung gốc có dấu hiệu mâu thuẫn, đã xử lý, **cần giáo viên xác nhận** |
| `[CẦN KIỂM CHỨNG]` | Số liệu **chưa đối chiếu được toàn văn văn bản gốc** |

> ### ⚠️ [LƯU Ý] ĐIỀU KIỆN SỬ DỤNG BẢN 1.0
>
> Bản 1.0 còn **49 mục mang nhãn `[CẦN KIỂM CHỨNG]`** — chủ yếu là số km/giờ chi tiết của chương trình đào tạo và biểu điểm sát hạch, nằm trong phần **Phụ lục** của Thông tư 17/2026/TT-BXD và Thông tư 108/2026/TT-BCA.
>
> **Không dùng số liệu mang nhãn này để cam kết với học viên khi chưa xác minh.** Danh mục đầy đủ ở **Phụ lục A** cuối tài liệu.

## CĂN CỨ PHÁP LÝ ÁP DỤNG

| Văn bản | Nội dung | Hiệu lực |
|---|---|:---:|
| **Luật 36/2024/QH15** — Trật tự, an toàn giao thông đường bộ | Hệ hạng GPLX mới (B, C1, C…) | 01/01/2025 |
| **Nghị định 160/2024/NĐ-CP** | Hoạt động đào tạo và sát hạch lái xe | 01/01/2025 |
| **Thông tư 14/2025/TT-BXD** | Quy định về đào tạo lái xe | 01/9/2025 |
| **Thông tư 17/2026/TT-BXD** | Sửa đổi TT 14/2025 — giảm thời lượng sa hình | **01/7/2026** |
| **Thông tư 108/2026/TT-BCA** | Sát hạch, cấp GPLX — bỏ thi mô phỏng | **01/7/2026** |
"""
rnd.render(GUIDE, first_h1_break=False)

# ═══════════════ MỤC LỤC ═══════════════
p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
r = p.add_run('MỤC LỤC')
r.font.size = Pt(19); r.bold = True; r.font.name = FONT
r.font.color.rgb = RGBColor.from_string(NAVY)
hrule(doc, NAVY, 12, 2, 12)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run('Nhấn Ctrl + A rồi F9 để cập nhật số trang sau khi chỉnh sửa tài liệu.')
r.italic = True; r.font.size = Pt(9); r.font.name = FONT
r.font.color.rgb = RGBColor.from_string(GREY_TXT)

# ---- TOC: field + nội dung dựng sẵn (điền từ toc.json nếu đã đo được) ----
import json as _json
TOCLIST = []
_tocf = os.path.join(WORK,'toc.json')
if os.path.exists(_tocf):
    TOCLIST = _json.load(open(_tocf, encoding='utf8'))

ptoc = doc.add_paragraph()
r1 = ptoc.add_run(); r1._r.append(_el('w:fldChar', **{'w:fldCharType': 'begin'}))
r2 = ptoc.add_run()
_it = OxmlElement('w:instrText'); _it.set(qn('xml:space'), 'preserve')
_it.text = r' TOC \o "1-2" \h \z \u '
r2._r.append(_it)
r3 = ptoc.add_run(); r3._r.append(_el('w:fldChar', **{'w:fldCharType': 'separate'}))

TOC_ENTRIES = []   # điền sau khi dựng thân tài liệu

def _toc_par(level, text, page):
    tp = doc.add_paragraph()
    tp.paragraph_format.space_after = Pt(1 if level == 2 else 5)
    tp.paragraph_format.space_before = Pt(6 if level == 1 else 0)
    tp.paragraph_format.left_indent = Cm(0 if level == 1 else 0.7)
    pr = tp._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tb = _el('w:tab', **{'w:val': 'right', 'w:pos': '9355'})
    tb.set(qn('w:leader'), 'dot')
    tabs.append(tb); pr.append(tabs)
    rr = tp.add_run(text)
    rr.font.name = FONT
    rr.font.size = Pt(11 if level == 1 else 9.5)
    rr.bold = (level == 1)
    rr.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else INK)
    tp.add_run('\t')
    rn = tp.add_run(str(page) if page else '')
    rn.font.name = FONT
    rn.font.size = Pt(11 if level == 1 else 9.5)
    rn.bold = (level == 1)
    rn.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else GREY_TXT)
    return tp

_toc_anchor = ptoc

# ═══════════════ NỘI DUNG CHÍNH ═══════════════
FILES = ['02_phan1_chuong_trinh.md', '03_phan2_quy_trinh_thi.md', '04_phan3_dat.md',
         '05_bai1.md', '06_bai2.md', '07_bai3.md', '08_bai4.md',
         '09_bai5.md', '10_bai6.md', '11_bai7.md']

DROP = re.compile(r'^>\s*\*\*Bản nháp')

for fn in FILES:
    md = open(os.path.join(DRAFTS, fn), encoding='utf8').read()
    # bỏ khối metadata nội bộ ở đầu file (blockquote bắt đầu bằng "**Bản nháp")
    lines = md.split('\n')
    out, i = [], 0
    while i < len(lines):
        if DROP.match(lines[i]):
            while i < len(lines) and lines[i].strip().startswith('>'):
                i += 1
            continue
        out.append(lines[i]); i += 1
    rnd.render('\n'.join(out), first_h1_break=True)

# ═══════════════ PHỤ LỤC A ═══════════════
APPENDIX = open(os.path.join(DRAFTS, '12_ra_soat_cheo.md'), encoding='utf8').read()
# lấy mục D và E
def section_of(md, start, end):
    i = md.find(start)
    j = md.find(end, i + 1) if end else len(md)
    return md[i:j] if i >= 0 else ''

apx = "# PHỤ LỤC A — DANH MỤC NỘI DUNG CẦN XÁC MINH\n"
apx += "### Dành cho giáo viên và người quản lý hồ sơ\n\n"
apx += ("Bản 1.0 của giáo trình này còn một số nội dung chưa đối chiếu được với toàn văn "
        "phụ lục của các Thông tư. Bảng dưới đây liệt kê đầy đủ để giáo viên bổ sung "
        "trước khi phát hành cho học viên.\n\n")
apx += section_of(APPENDIX, '## D. DANH MỤC', '## E. DANH MỤC').replace('## D. DANH MỤC', '## D. DANH MỤC')
apx += "\n\n"
apx += section_of(APPENDIX, '## E. DANH MỤC', '## F. TÀI SẢN')
apx += "\n\n"
apx += section_of(APPENDIX, '## B. BẢNG THUẬT NGỮ', '### B.5.')
rnd.render(apx, first_h1_break=True)

# ═══════════════ HOÀN TẤT MỤC LỤC ═══════════════
# thu thập heading 1 & 2 theo thứ tự tài liệu
_headings = []
for _p in doc.paragraphs:
    _sn = _p.style.name
    _t = _p.text.strip()
    if not _t:
        continue
    if _sn == 'Heading 1':
        _headings.append((1, _t))
    elif _sn == 'Heading 2':
        _headings.append((2, _t))

_body_after = doc.element.body
_pars_created = []
for _i, (_lv, _t) in enumerate(_headings):
    _pg = ''
    if _i < len(TOCLIST) and TOCLIST[_i].get('text') == _t:
        _pg = TOCLIST[_i].get('page', '')
    else:
        for _e in TOCLIST:
            if _e.get('text') == _t:
                _pg = _e.get('page', ''); break
    _pars_created.append(_toc_par(_lv, _t, _pg))
_rend = doc.add_run_paragraph if False else None
_endp = doc.add_paragraph()
_re_ = _endp.add_run(); _re_._r.append(_el('w:fldChar', **{'w:fldCharType': 'end'}))

# di chuyển các đoạn TOC + fldChar end lên ngay sau _toc_anchor
_prev = _toc_anchor._p
for _tp in _pars_created + [_endp]:
    _prev.addnext(_tp._p)
    _prev = _tp._p

# ═══════════════ HEADER / FOOTER ═══════════════
sec.different_first_page_header_footer = True

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_after = Pt(2)
r = hp.add_run('HỒ SƠ BÀI GIẢNG THỰC HÀNH LÁI XE Ô TÔ  |  HẠNG B – C1')
r.font.size = Pt(8); r.font.name = FONT; r.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
pPr = hp._p.get_or_add_pPr()
b = OxmlElement('w:pBdr')
b.append(_el('w:bottom', **{'w:val': 'single', 'w:sz': 6, 'w:space': '2', 'w:color': GREY_LINE}))
pPr.append(b)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
pPr = fp._p.get_or_add_pPr()
b = OxmlElement('w:pBdr')
b.append(_el('w:top', **{'w:val': 'single', 'w:sz': 6, 'w:space': '4', 'w:color': GREY_LINE}))
pPr.append(b)
tabs = OxmlElement('w:tabs')
tabs.append(_el('w:tab', **{'w:val': 'center', 'w:pos': '4677'}))
tabs.append(_el('w:tab', **{'w:val': 'right', 'w:pos': '9355'}))
pPr.append(tabs)
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = fp.add_run('%s — %s' % (VERSION, UPDATED))
r.font.size = Pt(8); r.font.name = FONT
r.font.color.rgb = RGBColor.from_string(GREY_TXT)
fp.add_run('\t')
r = fp.add_run('Trang ')
r.font.size = Pt(8.5); r.font.name = FONT; r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
for rr in field(fp, r' PAGE \* MERGEFORMAT '):
    rr.font.size = Pt(8.5); rr.font.name = FONT; rr.bold = True
    rr.font.color.rgb = RGBColor.from_string(NAVY)
r = fp.add_run(' / ')
r.font.size = Pt(8.5); r.font.name = FONT
r.font.color.rgb = RGBColor.from_string(GREY_TXT)
for rr in field(fp, r' NUMPAGES \* MERGEFORMAT '):
    rr.font.size = Pt(8.5); rr.font.name = FONT
    rr.font.color.rgb = RGBColor.from_string(GREY_TXT)
fp.add_run('\t')
r = fp.add_run('Đào tạo lái xe ô tô')
r.font.size = Pt(8); r.font.name = FONT
r.font.color.rgb = RGBColor.from_string(GREY_TXT)

# trang bìa: không header/footer
for para in list(sec.first_page_header.paragraphs) + list(sec.first_page_footer.paragraphs):
    for rr in list(para.runs):
        rr._r.getparent().remove(rr._r)

# cập nhật field khi mở
sp = doc.settings.element
u = OxmlElement('w:updateFields'); u.set(qn('w:val'), 'true')
sp.append(u)

doc.save(OUT)
print('ĐÃ TẠO:', OUT, os.path.getsize(OUT) // 1024, 'KB')
